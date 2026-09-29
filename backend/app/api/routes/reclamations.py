from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import Response, FileResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func, or_
from sqlalchemy.orm import selectinload
from datetime import date, timedelta
import secrets
from collections import defaultdict
from io import BytesIO
from app.core.database import get_db
from app.core.deps import get_current_user, require_role
from app.models.user import User, Role, RoleEnum
from app.models.contribuable import Contribuable
from app.models.reclamation import (
    Reclamation, StatutReclamationEnum,
    TypeReclamation, MotifReclamation, ActionHistorique, TypeActionEnum,
    Affectation, RoleDossierEnum,
)
from app.schemas.reclamation import (
    ReclamationCreate, ReclamationRead, ReclamationList,
    QualifierReclamation, ReclamationUpdate,
    TypeReclamationRead, MotifReclamationRead,
    ActionHistoriqueRead, DemanderPieces,
    AffectationAssign, AffectationRead,
)
from app.services import pdf_service
from app.services import minio_service
from app.core.config import settings

router = APIRouter()


@router.get("/types", response_model=list[TypeReclamationRead])
async def list_types_reclamation(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    result = await db.execute(
        select(TypeReclamation)
        .where(TypeReclamation.actif.is_(True))
        .order_by(TypeReclamation.id)
    )
    return result.scalars().all()


@router.get("/motifs", response_model=list[MotifReclamationRead])
async def list_motifs_reclamation(
    id_type: int | None = None,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    query = select(MotifReclamation)
    if id_type:
        query = query.where(MotifReclamation.id_type == id_type)
    result = await db.execute(query.order_by(MotifReclamation.id))
    return result.scalars().all()


@router.get("/", response_model=ReclamationList)
async def list_reclamations(
    page: int = Query(1, ge=1),
    size: int = Query(20, ge=1, le=100),
    statut: str | None = None,
    id_type: int | None = None,
    search: str | None = None,
    date_debut: date | None = None,
    date_fin: date | None = None,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role("ADMIN", "SAISIE", "INSTRUCTEUR", "CHEF", "DIRECTEUR")),
):
    query = select(Reclamation).options(
        selectinload(Reclamation.contribuable),
        selectinload(Reclamation.type),
        selectinload(Reclamation.motif),
    )
    count_query = select(func.count(Reclamation.id))

    if statut:
        query = query.where(Reclamation.statut == statut)
        count_query = count_query.where(Reclamation.statut == statut)
    if id_type:
        query = query.where(Reclamation.id_type == id_type)
        count_query = count_query.where(Reclamation.id_type == id_type)
    if search:
        filtre = or_(
            Reclamation.numero_dossier.ilike(f"%{search}%"),
            Reclamation.reference_imposition.ilike(f"%{search}%"),
            Reclamation.code_suivi.ilike(f"%{search}%"),
            Reclamation.contribuable.has(
                or_(
                    Contribuable.numero_fiscal.ilike(f"%{search}%"),
                    Contribuable.nom_raison_sociale.ilike(f"%{search}%"),
                )
            ),
        )
        query = query.where(filtre)
        count_query = count_query.where(filtre)
    if date_debut:
        query = query.where(Reclamation.date_depot >= date_debut)
        count_query = count_query.where(Reclamation.date_depot >= date_debut)
    if date_fin:
        query = query.where(Reclamation.date_depot <= date_fin)
        count_query = count_query.where(Reclamation.date_depot <= date_fin)

    total_result = await db.execute(count_query)
    total = total_result.scalar()

    query = query.offset((page - 1) * size).limit(size).order_by(Reclamation.created_at.desc())
    result = await db.execute(query)
    items = result.scalars().all()

    return ReclamationList(
        items=[ReclamationRead.model_validate(r) for r in items],
        total=total,
        page=page,
        size=size,
    )


@router.post("/", response_model=ReclamationRead, status_code=status.HTTP_201_CREATED)
async def create_reclamation(
    data: ReclamationCreate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role("ADMIN", "SAISIE", "INSTRUCTEUR")),
):
    result = await db.execute(
        select(Reclamation).order_by(Reclamation.id.desc()).limit(1)
    )
    last = result.scalar_one_or_none()
    year = date.today().year
    if last and last.numero_dossier.startswith(f"REC-{year}"):
        num = int(last.numero_dossier.split("-")[-1]) + 1
    else:
        num = 1
    numero_dossier = f"REC-{year}-{num:06d}"

    reclamation = Reclamation(
        numero_dossier=numero_dossier,
        code_suivi=_generer_code_suivi(),
        id_contribuable=data.id_contribuable,
        id_type=data.id_type,
        id_motif=data.id_motif,
        canal_entree=data.canal_entree,
        resume_faits=data.resume_faits,
        montant_concerne=data.montant_concerne,
        reference_imposition=data.reference_imposition,
        id_agent_createur=current_user.id,
        statut=StatutReclamationEnum.ENREGISTREE,
    )
    db.add(reclamation)
    await db.flush()

    action = ActionHistorique(
        id_reclamation=reclamation.id,
        id_agent=current_user.id,
        action=TypeActionEnum.CREATION,
        commentaire="Reclamation creee",
    )
    db.add(action)

    try:
        full = (await db.execute(
            select(Reclamation)
            .options(selectinload(Reclamation.contribuable), selectinload(Reclamation.type))
            .where(Reclamation.id == reclamation.id)
        )).scalar_one()
        reclamation.pdf_accuse_path = pdf_service.generer_et_stocker_accuse(full)
    except Exception as exc:
        reclamation.pdf_accuse_path = None
        print(f"[WARN] Generation de l'accuse impossible: {exc}")

    await db.flush()

    return ReclamationRead.model_validate(reclamation)


@router.get("/{reclamation_id}/accuse")
async def telecharger_accuse(
    reclamation_id: int,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    result = await db.execute(
        select(Reclamation).where(Reclamation.id == reclamation_id)
    )
    reclamation = result.scalar_one_or_none()
    if not reclamation:
        raise HTTPException(status_code=404, detail="Reclamation introuvable")
    if not reclamation.pdf_accuse_path:
        raise HTTPException(status_code=404, detail="Accuse de reception non genere")

    chemin = reclamation.pdf_accuse_path
    if chemin.startswith("local:"):
        filename = chemin.split(":", 1)[1]
        filepath = minio_service.get_local_file(filename)
        if not filepath.exists():
            raise HTTPException(status_code=404, detail="Fichier absent du disque")
        return FileResponse(
            path=str(filepath),
            filename=f"accuse_{reclamation.numero_dossier}.pdf",
            media_type="application/pdf",
        )

    return {"url": minio_service.presigned_download(chemin)}


TERMINAUX = ["CLOTUREE", "REJETEE", "NOTIFIEE"]


def _statut_texte(v) -> str:
    return v.value if hasattr(v, "value") else str(v)


def _generer_code_suivi() -> str:
    return f"{date.today().strftime('%Y%m')}-{secrets.token_hex(3).upper()}"


@router.get("/stats")
async def stats_reclamations(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role("ADMIN", "CHEF", "DIRECTEUR", "INSTRUCTEUR", "SAISIE")),
):
    today = date.today()
    six_mois = (today.replace(day=1) - timedelta(days=150)).replace(day=1)

    total = await db.scalar(select(func.count(Reclamation.id))) or 0

    rows_statut = await db.execute(
        select(Reclamation.statut, func.count(Reclamation.id)).group_by(Reclamation.statut)
    )
    par_statut = {str(r[0].value if hasattr(r[0], 'value') else r[0]): r[1] for r in rows_statut}

    rows_type = await db.execute(
        select(TypeReclamation.libelle, func.count(Reclamation.id))
        .outerjoin(Reclamation, Reclamation.id_type == TypeReclamation.id)
        .group_by(TypeReclamation.libelle)
    )
    par_type = [{"libelle": r[0], "total": r[1]} for r in rows_type]

    rows_canal = await db.execute(
        select(Reclamation.canal_entree, func.count(Reclamation.id))
        .group_by(Reclamation.canal_entree)
    )
    par_canal = {str(r[0].value if hasattr(r[0], 'value') else r[0]): r[1] for r in rows_canal}

    rows_mois = await db.execute(
        select(
            func.to_char(Reclamation.date_depot, 'YYYY-MM').label('mois'),
            func.count(Reclamation.id),
        )
        .where(Reclamation.date_depot >= six_mois)
        .group_by('mois')
        .order_by('mois')
    )
    par_mois = [{"mois": r[0], "total": r[1]} for r in rows_mois]

    en_retard = await db.scalar(
        select(func.count(Reclamation.id)).where(
            Reclamation.statut.not_in(TERMINAUX),
            Reclamation.date_limite_reponse < today,
        )
    ) or 0

    en_cours = await db.scalar(
        select(func.count(Reclamation.id)).where(
            Reclamation.statut.not_in(TERMINAUX),
        )
    ) or 0

    cloturees_rows = await db.execute(
        select(Reclamation.date_depot).where(Reclamation.statut.in_(["CLOTUREE", "NOTIFIEE"]))
    )
    delais = [(today - r[0]).days for r in cloturees_rows if r[0]]
    delai_moyen = round(sum(delais) / len(delais), 1) if delais else 0

    rows_agent = await db.execute(
        select(
            User.nom,
            User.id,
            func.count(Reclamation.id).label('traitees'),
        )
        .outerjoin(Reclamation, Reclamation.id_agent_createur == User.id)
        .group_by(User.id, User.nom)
        .having(func.count(Reclamation.id) > 0)
    )
    par_agent = [{"nom": r[0], "traitees": r[2]} for r in rows_agent]

    return {
        "total": total,
        "en_cours": en_cours,
        "en_retard": en_retard,
        "delai_moyen_jours": delai_moyen,
        "par_statut": par_statut,
        "par_type": par_type,
        "par_canal": par_canal,
        "par_mois": par_mois,
        "par_agent": par_agent,
    }


@router.get("/export")
async def export_reclamations(
    format: str = Query("xlsx", pattern="^(xlsx|csv)$"),
    statut: str | None = None,
    id_type: int | None = None,
    date_debut: date | None = None,
    date_fin: date | None = None,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role("ADMIN", "CHEF", "DIRECTEUR", "INSTRUCTEUR")),
):
    today = date.today()

    query = select(Reclamation).options(
        selectinload(Reclamation.contribuable),
        selectinload(Reclamation.type),
    )
    if statut:
        query = query.where(Reclamation.statut == statut)
    if id_type:
        query = query.where(Reclamation.id_type == id_type)
    if date_debut:
        query = query.where(Reclamation.date_depot >= date_debut)
    if date_fin:
        query = query.where(Reclamation.date_depot <= date_fin)
    query = query.order_by(Reclamation.created_at.desc())
    result = await db.execute(query)
    rows = result.scalars().all()

    STATUT_FR = {
        "ENREGISTREE": "Enregistree",
        "A_QUALIFIER": "A qualifier",
        "EN_INSTRUCTION": "En instruction",
        "EN_ATTENTE_PIECES": "En attente pieces",
        "PROJET_REPONSE": "Projet reponse",
        "EN_VALIDATION": "En validation",
        "EN_VISA_DIRECTEUR": "Visa directeur",
        "SIGNEE": "Signee",
        "NOTIFIEE": "Notifiee",
        "CLOTUREE": "Cloturee",
        "REJETEE": "Rejetee",
        "CONTENTIEUX_JUDICIAIRE": "Contentieux judiciaire",
    }

    def _row(r):
        s = str(r.statut.value if hasattr(r.statut, 'value') else r.statut)
        return [
            r.numero_dossier,
            r.contribuable.nom_raison_sociale if r.contribuable else '',
            r.type.libelle if r.type else '',
            s,
            str(r.date_depot) if r.date_depot else '',
            str(r.date_limite_reponse) if r.date_limite_reponse else '',
            STATUT_FR.get(s, s),
            f"{r.montant_concerne:,.0f} DA" if r.montant_concerne else '',
            r.reference_imposition or '',
            r.canal_entree.value if hasattr(r.canal_entree, 'value') else str(r.canal_entree),
        ]

    header = [
        "Numero dossier", "Contribuable", "Type", "Statut", "Date depot",
        "Date limite", "Statut FR", "Montant", "Ref imposition", "Canal",
    ]

    if format == "csv":
        import csv
        import io
        buf = BytesIO()
        tw = io.TextIOWrapper(buf, encoding="utf-8", newline="")
        writer = csv.writer(tw, delimiter=";")
        writer.writerow(header)
        for r in rows:
            writer.writerow(_row(r))
        tw.flush()
        content = buf.getvalue()
        tw.detach()
        return Response(
            content=content,
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename=reclamations_{today}.csv"},
        )

    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reclamations"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1a5319", fill_type="solid")
    thin = Side(style="thin", color="CCCCCC")
    border = Border(top=thin, bottom=thin, left=thin, right=thin)

    for col, label in enumerate(header, 1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    for i, r in enumerate(rows, 2):
        for j, val in enumerate(_row(r), 1):
            cell = ws.cell(row=i, column=j, value=val)
            cell.border = border

    for col in range(1, len(header) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 18

    buf = BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=reclamations_{today}.xlsx"},
    )


@router.get("/a-qualifier")
async def reclamations_a_qualifier(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role("ADMIN", "SAISIE")),
):
    """File d'attente de qualification pour l'agent de saisie.

    Retourne les dossiers nouvellement reçus (ENREGISTREE / A_QUALIFIER)
    qui doivent etre categories et transmis a l'instruction.
    """
    cibles = [
        StatutReclamationEnum.ENREGISTREE,
        StatutReclamationEnum.A_QUALIFIER,
    ]

    query = (
        select(Reclamation)
        .options(
            selectinload(Reclamation.contribuable),
            selectinload(Reclamation.type),
            selectinload(Reclamation.motif),
        )
        .where(Reclamation.statut.in_(cibles))
        .order_by(Reclamation.created_at.desc())
    )
    result = await db.execute(query)
    items = result.scalars().all()

    def key_statut(s) -> str:
        return s.value if hasattr(s, "value") else str(s)

    dossiers = []
    for r in items:
        dossiers.append(
            {
                "id": r.id,
                "numero_dossier": r.numero_dossier,
                "statut": key_statut(r.statut),
                "canal_entree": key_statut(r.canal_entree),
                "date_depot": r.date_depot.isoformat() if r.date_depot else None,
                "montant_concerne": float(r.montant_concerne) if r.montant_concerne is not None else None,
                "reference_imposition": r.reference_imposition,
                "resume_faits": r.resume_faits,
                "contribuable": {
                    "nom_raison_sociale": r.contribuable.nom_raison_sociale if r.contribuable else None,
                    "numero_fiscal": r.contribuable.numero_fiscal if r.contribuable else None,
                },
            }
        )

    return {"total": len(dossiers), "items": dossiers}


@router.get("/a-instruire")
async def reclamations_a_instruire(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role("ADMIN", "INSTRUCTEUR")),
):
    """File d'attente d'instruction pour le role instructeur.

    Retourne les dossiers en EN_INSTRUCTION / EN_ATTENTE_PIECES / PROJET_REPONSE
    (circuit instructeur), avec le nombre de pieces et l'etat de la decision.
    """
    cibles = [
        StatutReclamationEnum.EN_INSTRUCTION,
        StatutReclamationEnum.EN_ATTENTE_PIECES,
        StatutReclamationEnum.PROJET_REPONSE,
    ]

    query = (
        select(Reclamation)
        .options(
            selectinload(Reclamation.contribuable),
            selectinload(Reclamation.type),
            selectinload(Reclamation.motif),
            selectinload(Reclamation.decision),
            selectinload(Reclamation.pieces),
            selectinload(Reclamation.affectations).selectinload(Affectation.agent),
        )
        .where(Reclamation.statut.in_(cibles))
        .order_by(Reclamation.created_at.desc())
    )
    result = await db.execute(query)
    items = result.scalars().all()

    def key_statut(s) -> str:
        return s.value if hasattr(s, "value") else str(s)

    dossiers = []
    for r in items:
        decision = r.decision
        affect_actifs = [a for a in (r.affectations or []) if a.actif and a.role_dossier == RoleDossierEnum.INSTRUCTEUR]
        dossiers.append(
            {
                "id": r.id,
                "numero_dossier": r.numero_dossier,
                "statut": key_statut(r.statut),
                "date_depot": r.date_depot.isoformat() if r.date_depot else None,
                "date_limite_reponse": r.date_limite_reponse.isoformat() if r.date_limite_reponse else None,
                "montant_concerne": float(r.montant_concerne) if r.montant_concerne is not None else None,
                "reference_imposition": r.reference_imposition,
                "resume_faits": r.resume_faits,
                "nb_pieces": len(r.pieces) if r.pieces else 0,
                "instructeur": {
                    "id_agent": affect_actifs[0].agent.id,
                    "nom": affect_actifs[0].agent.nom,
                } if affect_actifs and affect_actifs[0].agent else None,
                "contribuable": {
                    "nom_raison_sociale": r.contribuable.nom_raison_sociale if r.contribuable else None,
                    "numero_fiscal": r.contribuable.numero_fiscal if r.contribuable else None,
                },
                "type": {"libelle": r.type.libelle if r.type else None, "code": key_statut(r.type.code) if r.type else None},
                "motif": {"libelle": r.motif.libelle if r.motif else None},
                "decision": None if not decision else {
                    "existe": True,
                    "statut": key_statut(decision.statut),
                    "type_decision": key_statut(decision.type_decision),
                },
            }
        )

    return {"total": len(dossiers), "items": dossiers}


@router.get("/a-valider")
async def reclamations_a_valider(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role("ADMIN", "CHEF", "DIRECTEUR")),
):
    """File d'attente de validation selon le role de l'utilisateur connecte.

    - CHEF : decisions en attente d'avis (statut EN_VALIDATION)
    - DIRECTEUR : decisions en attente de visa/signature (EN_VISA_DIRECTEUR)
    - ADMIN : voit les deux files.
    """
    role = current_user.role.libelle

    stats_statuts = ["EN_VALIDATION", "EN_VISA_DIRECTEUR"]

    def statut_cible(role: str) -> list[StatutReclamationEnum]:
        if role == "DIRECTEUR":
            return [StatutReclamationEnum.EN_VISA_DIRECTEUR]
        if role == "CHEF":
            return [StatutReclamationEnum.EN_VALIDATION]
        return [StatutReclamationEnum.EN_VALIDATION, StatutReclamationEnum.EN_VISA_DIRECTEUR]

    cibles = statut_cible(role)

    def key_statut(s: StatutReclamationEnum) -> str:
        return s.value if hasattr(s, "value") else str(s)

    # Compteurs par statut (pour les deux files affichables)
    rows = await db.execute(
        select(Reclamation.statut, func.count(Reclamation.id))
        .where(Reclamation.statut.in_([StatutReclamationEnum(s) for s in stats_statuts]))
        .group_by(Reclamation.statut)
    )
    compteurs = {}
    for r in rows:
        label = key_statut(r[0])
        compteurs[label] = r[1]

    # Dossiers a valider pour ce role
    query = (
        select(Reclamation)
        .options(
            selectinload(Reclamation.contribuable),
            selectinload(Reclamation.type),
            selectinload(Reclamation.motif),
            selectinload(Reclamation.decision),
        )
        .where(Reclamation.statut.in_(cibles))
        .order_by(Reclamation.date_limite_reponse.asc().nulls_last(), Reclamation.created_at.asc())
    )
    items_result = await db.execute(query)
    items = items_result.scalars().all()

    dossiers = []
    for r in items:
        decision = r.decision
        dossiers.append(
            {
                "id": r.id,
                "numero_dossier": r.numero_dossier,
                "statut": key_statut(r.statut),
                "date_depot": r.date_depot.isoformat() if r.date_depot else None,
                "date_limite_reponse": r.date_limite_reponse.isoformat() if r.date_limite_reponse else None,
                "montant_concerne": float(r.montant_concerne) if r.montant_concerne is not None else None,
                "contribuable": {
                    "nom_raison_sociale": r.contribuable.nom_raison_sociale if r.contribuable else None,
                    "numero_fiscal": r.contribuable.numero_fiscal if r.contribuable else None,
                },
                "type": {"libelle": r.type.libelle if r.type else None, "code": r.type.code.value if r.type and hasattr(r.type.code, "value") else (r.type.code if r.type else None)},
                "decision": {
                    "id": decision.id if decision else None,
                    "type_decision": key_statut(decision.type_decision) if decision else None,
                    "fondement_juridique": decision.fondement_juridique if decision else None,
                    "montant_accorde": float(decision.montant_accorde) if decision else None,
                    "montant_rejete": float(decision.montant_rejete) if decision else None,
                },
            }
        )

    return {
        "role": role,
        "compteurs": {
            "en_validation": compteurs.get("EN_VALIDATION", 0),
            "en_visa_directeur": compteurs.get("EN_VISA_DIRECTEUR", 0),
        },
        "total": len(dossiers),
        "items": dossiers,
    }


@router.get("/a-relancer")
async def reclamations_a_relancer(
    jours: int = Query(15, ge=1, le=90),
    inclu_sans_limite: bool = Query(False),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role("ADMIN", "CHEF", "DIRECTEUR", "INSTRUCTEUR")),
):
    """Reclamations non cloturees dont la date limite tombe sous X jours (relances).

    Retourne egalement les dossiers deja en retard (date limite passee) afin de
    relancer par priorite croissante.
    """
    aujourd = date.today()
    horizon = aujourd + timedelta(days=jours)

    def key_statut(s: StatutReclamationEnum) -> str:
        return s.value if hasattr(s, "value") else str(s)

    conditions = [
        Reclamation.statut.not_in(TERMINAUX),
    ]
    if not inclu_sans_limite:
        conditions.append(Reclamation.date_limite_reponse.isnot(None))

    query = (
        select(Reclamation)
        .options(
            selectinload(Reclamation.contribuable),
            selectinload(Reclamation.type),
            selectinload(Reclamation.motif),
            selectinload(Reclamation.affectations),
        )
        .where(*conditions)
        .where(
            or_(
                Reclamation.date_limite_reponse < aujourd,
                Reclamation.date_limite_reponse <= horizon,
            )
        )
        .order_by(Reclamation.date_limite_reponse.asc().nulls_last())
    )
    result = await db.execute(query)
    items = result.scalars().all()

    instructeurs = {}
    for a in await db.execute(
        select(Affectation, User.nom).outerjoin(User, User.id == Affectation.id_agent).where(Affectation.actif.is_(True))
    ):
        affectation, nom = a[0], a[1]
        instructeurs[affectation.id_reclamation] = {
            "id_agent": affectation.id_agent,
            "nom": nom,
            "role_dossier": key_statut(affectation.role_dossier),
        }

    dossiers = []
    for r in items:
        limite = r.date_limite_reponse
        jours_restants = (limite - aujourd).days if limite else None
        dossiers.append(
            {
                "id": r.id,
                "numero_dossier": r.numero_dossier,
                "statut": key_statut(r.statut),
                "canal_entree": key_statut(r.canal_entree),
                "date_depot": r.date_depot.isoformat() if r.date_depot else None,
                "date_limite_reponse": limite.isoformat() if limite else None,
                "jours_restants": jours_restants,
                "en_retard": jours_restants is not None and jours_restants < 0,
                "montant_concerne": float(r.montant_concerne) if r.montant_concerne is not None else None,
                "reference_imposition": r.reference_imposition,
                "contribuable": {
                    "nom_raison_sociale": r.contribuable.nom_raison_sociale if r.contribuable else None,
                    "numero_fiscal": r.contribuable.numero_fiscal if r.contribuable else None,
                },
                "type": {"libelle": r.type.libelle if r.type else None, "code": key_statut(r.type.code) if r.type else None},
                "instructeur_actuel": instructeurs.get(r.id),
            }
        )

    return {
        "items": dossiers,
        "total": len(dossiers),
        "page": 1,
        "size": len(dossiers),
        "horizon_jours": jours,
    }


@router.get("/{reclamation_id}/pdf")
async def exporter_dossier_pdf(
    reclamation_id: int,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role("ADMIN", "SAISIE", "INSTRUCTEUR", "CHEF", "DIRECTEUR")),
):
    """Exporte le dossier complet (frise, historique, pieces, decision) en PDF."""
    result = await db.execute(
        select(Reclamation)
        .options(
            selectinload(Reclamation.contribuable),
            selectinload(Reclamation.type),
            selectinload(Reclamation.motif),
            selectinload(Reclamation.historique),
            selectinload(Reclamation.pieces),
            selectinload(Reclamation.decision),
            selectinload(Reclamation.affectations),
        )
        .where(Reclamation.id == reclamation_id)
    )
    reclamation = result.scalar_one_or_none()
    if not reclamation:
        raise HTTPException(status_code=404, detail="Reclamation introuvable")

    try:
        pdf_bytes = pdf_service.build_dossier_pdf(reclamation)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Generation du PDF impossible: {exc}")

    filename = f"dossier_{reclamation.numero_dossier}.pdf"
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{filename}",
        },
    )


@router.get("/{reclamation_id}/historique", response_model=list[ActionHistoriqueRead])
async def get_historique(
    reclamation_id: int,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    result = await db.execute(select(Reclamation).where(Reclamation.id == reclamation_id))
    if not result.scalar_one_or_none():
        raise HTTPException(status_code=404, detail="Reclamation introuvable")

    actions = await db.execute(
        select(ActionHistorique)
        .where(ActionHistorique.id_reclamation == reclamation_id)
        .order_by(ActionHistorique.created_at.asc())
    )
    return actions.scalars().all()


@router.post("/{reclamation_id}/demander-pieces", response_model=ReclamationRead)
async def demander_pieces(
    reclamation_id: int,
    data: DemanderPieces,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role("ADMIN", "INSTRUCTEUR")),
):
    result = await db.execute(select(Reclamation).where(Reclamation.id == reclamation_id))
    reclamation = result.scalar_one_or_none()
    if not reclamation:
        raise HTTPException(status_code=404, detail="Reclamation introuvable")

    if reclamation.statut in (
        StatutReclamationEnum.CLOTUREE,
        StatutReclamationEnum.REJETEE,
        StatutReclamationEnum.CONTENTIEUX_JUDICIAIRE,
    ):
        raise HTTPException(status_code=400, detail=f"Statut {reclamation.statut.value} clos")

    reclamation.statut = StatutReclamationEnum.EN_ATTENTE_PIECES

    action = ActionHistorique(
        id_reclamation=reclamation_id,
        id_agent=current_user.id,
        action=TypeActionEnum.DEMANDE_PIECES,
        commentaire=data.motif,
        donnees_contexte={"pieces": data.liste_pieces},
    )
    db.add(action)
    await db.flush()

    return ReclamationRead.model_validate(reclamation)


@router.get("/instructeurs")
async def lister_instructeurs(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role("ADMIN", "CHEF", "DIRECTEUR", "INSTRUCTEUR")),
):
    """Liste des instructeurs actifs pour l'affectation des dossiers."""
    rows = await db.execute(
        select(User)
        .join(User.role)
        .where(Role.libelle == RoleEnum.INSTRUCTEUR, User.actif.is_(True))
        .order_by(User.nom)
    )
    return [
        {"id": u.id, "nom": u.nom, "email": u.email}
        for u in rows.scalars().all()
    ]


@router.get("/{reclamation_id}/affectations", response_model=list[AffectationRead])
async def lister_affectations(
    reclamation_id: int,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role("ADMIN", "INSTRUCTEUR", "CHEF", "DIRECTEUR")),
):
    result = await db.execute(
        select(Affectation, User.nom)
        .join(User, User.id == Affectation.id_agent)
        .where(Affectation.id_reclamation == reclamation_id)
        .order_by(Affectation.date_debut.desc())
    )
    items = []
    for r in result.all():
        affectation, nom_agent = r
        items.append(
            AffectationRead(
                id=affectation.id,
                id_agent=affectation.id_agent,
                nom_agent=nom_agent,
                role_dossier=_statut_texte(affectation.role_dossier),
                date_debut=affectation.date_debut,
                date_fin=affectation.date_fin,
                actif=affectation.actif,
            )
        )
    return items


@router.post("/{reclamation_id}/affectation", response_model=AffectationRead)
async def affecter_reclamation(
    reclamation_id: int,
    data: AffectationAssign,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role("ADMIN", "CHEF", "DIRECTEUR", "INSTRUCTEUR")),
):
    """Affecte un agent instructeur (ou un autre role) au dossier.

    L'affectation precedente du meme role est cloturee (actif=False, date_fin).
    """
    result = await db.execute(select(Reclamation).where(Reclamation.id == reclamation_id))
    reclamation = result.scalar_one_or_none()
    if not reclamation:
        raise HTTPException(status_code=404, detail="Reclamation introuvable")

    agent_rows = await db.execute(select(User).where(User.id == data.id_agent))
    agent = agent_rows.scalar_one_or_none()
    if not agent:
        raise HTTPException(status_code=400, detail="Agent introuvable")

    role_dossier = RoleDossierEnum(data.role)
    aujourd = date.today()

    actifs = await db.execute(
        select(Affectation).where(
            Affectation.id_reclamation == reclamation_id,
            Affectation.role_dossier == role_dossier,
            Affectation.actif.is_(True),
        )
    )
    for affect in actifs.scalars().all():
        affect.actif = False
        affect.date_fin = aujourd

    affectation = Affectation(
        id_reclamation=reclamation_id,
        id_agent=agent.id,
        role_dossier=role_dossier,
        date_debut=aujourd,
        actif=True,
    )
    db.add(affectation)
    await db.flush()

    if reclamation.statut == StatutReclamationEnum.A_QUALIFIER:
        reclamation.statut = StatutReclamationEnum.EN_INSTRUCTION

    action = ActionHistorique(
        id_reclamation=reclamation_id,
        id_agent=current_user.id,
        action=TypeActionEnum.AFFECTATION,
        commentaire=f"Affectation: {agent.nom} ({role_dossier.value})",
    )
    db.add(action)

    return AffectationRead(
        id=affectation.id,
        id_agent=agent.id,
        nom_agent=agent.nom,
        role_dossier=_statut_texte(role_dossier),
        date_debut=affectation.date_debut,
        date_fin=affectation.date_fin,
        actif=affectation.actif,
    )


@router.delete("/{reclamation_id}/affectation/{affectation_id}")
async def desaffecter_reclamation(
    reclamation_id: int,
    affectation_id: int,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role("ADMIN", "CHEF", "DIRECTEUR", "INSTRUCTEUR")),
):
    result = await db.execute(
        select(Affectation).where(
            Affectation.id == affectation_id,
            Affectation.id_reclamation == reclamation_id,
        )
    )
    affectation = result.scalar_one_or_none()
    if not affectation:
        raise HTTPException(status_code=404, detail="Affectation introuvable")

    affectation.actif = False
    affectation.date_fin = date.today()
    action = ActionHistorique(
        id_reclamation=reclamation_id,
        id_agent=current_user.id,
        action=TypeActionEnum.AFFECTATION,
        commentaire=f"Desaffectation de l'agent n{affectation.id_agent}",
    )
    db.add(action)
    await db.flush()
    return {"detail": "Agent desaffecte"}


@router.get("/{reclamation_id}", response_model=ReclamationRead)
async def get_reclamation(
    reclamation_id: int,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    result = await db.execute(
        select(Reclamation).where(Reclamation.id == reclamation_id)
    )
    reclamation = result.scalar_one_or_none()

    if not reclamation:
        raise HTTPException(status_code=404, detail="Reclamation introuvable")

    return ReclamationRead.model_validate(reclamation)


@router.patch("/{reclamation_id}/qualifier", response_model=ReclamationRead)
async def qualifier_reclamation(
    reclamation_id: int,
    data: QualifierReclamation,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role("ADMIN", "SAISIE", "INSTRUCTEUR")),
):
    result = await db.execute(
        select(Reclamation).where(Reclamation.id == reclamation_id)
    )
    reclamation = result.scalar_one_or_none()

    if not reclamation:
        raise HTTPException(status_code=404, detail="Reclamation introuvable")

    if reclamation.statut not in (StatutReclamationEnum.ENREGISTREE, StatutReclamationEnum.A_QUALIFIER):
        raise HTTPException(
            status_code=400,
            detail=f"Statut actuel {reclamation.statut} ne permet pas la qualification",
        )

    type_result = await db.execute(
        select(TypeReclamation).where(TypeReclamation.id == data.id_type)
    )
    type_recl = type_result.scalar_one_or_none()
    if not type_recl:
        raise HTTPException(status_code=400, detail="Type de reclamation introuvable")

    reclamation.id_type = data.id_type
    reclamation.id_motif = data.id_motif
    reclamation.date_limite_reponse = date.today() + timedelta(days=type_recl.delai_legal_jours)
    reclamation.statut = StatutReclamationEnum.EN_INSTRUCTION

    action = ActionHistorique(
        id_reclamation=reclamation.id,
        id_agent=current_user.id,
        action=TypeActionEnum.QUALIFICATION,
        commentaire=f"Type: {type_recl.code}, Date limite: {reclamation.date_limite_reponse}",
    )
    db.add(action)

    return ReclamationRead.model_validate(reclamation)
